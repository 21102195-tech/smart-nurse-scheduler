"""스마트 널스 스케줄러 v2 — 관리자용 초안 생성 및 독립 검수.
실행: python -m streamlit run smart_nurse_scheduler_v2.py --server.address 127.0.0.1
핵심 함수는 Streamlit 없이 import하여 테스트할 수 있습니다.
"""
from __future__ import annotations
import calendar
import copy
import hashlib
import io
import math
import random
import re
import time
from dataclasses import dataclass, asdict
from datetime import date
from typing import Callable
import numpy as np
import pandas as pd

SHIFTS = ('D', 'E', 'N', 'DE', 'OFF', '교육')
DUTIES = ('D', 'E', 'N', 'DE')
ALIASES = {
    'D':'D', '데이':'D', 'DAY':'D', 'E':'E', '이브':'E', '이브닝':'E', 'EVENING':'E',
    'N':'N', 'NC1':'N', '나이트':'N', 'NIGHT':'N', 'DE':'DE',
    'OFF':'OFF', '오프':'OFF', '휴무':'OFF', '휴':'OFF',
    **{x:'OFF' for x in ['C','OF','OF1','V','H','H1','NV','BV','B','S','S10','S2','S3']},
    **{x:'교육' for x in ['CPE','P','PE','PE1','PL','교육','EDU']},
}
FORBIDDEN = (
    ('D','D','N','N','N'), ('D','D','D','N','N'), ('D','D','D','D','N'),
    ('D','E','N','N','N'), ('E','E','N','N','N'), ('D','D','E','N','N'),
    ('DE','DE','N','N','N'), ('DE','DE','DE','N','N'),
    ('DE','DE','DE','DE','N'), ('DE','E','N','N','N'), ('DE','DE','E','N','N'),
)
ISSUE_COLUMNS = ['구분','직원','날짜','항목','내용']

@dataclass(frozen=True)
class Settings:
    max_work: int = 5
    max_night: int = 6
    max_consecutive_night: int = 3
    night_keeper_target: int = 15
    two_off_after_five: bool = True
    no_single_night: bool = True
    group_balance: bool = True
    two_off_after_night: bool = True
    no_single_work: bool = True  # 희망조건: 필수조건과 구분
    education_counts_as_day: bool = False
    max_iterations: int = 60000
    time_limit: float = 30.0
    seed: int = 42

@dataclass
class Problem:
    source: pd.DataFrame
    year: int
    month: int
    names: list[str]
    groups: list[str]
    keepers: list[bool]
    allowed: list[set[str]]
    wanted: list[set[int]]
    fixed: np.ndarray
    fixed_codes: np.ndarray
    histories: list[list[str]]
    warnings: list[dict]
    requirements: dict[str, list[int]]
    row_indices: list[int]

    @property
    def days(self): return calendar.monthrange(self.year, self.month)[1]

    @property
    def dates(self): return [date(self.year, self.month, d) for d in range(1, self.days+1)]


def blank(value):
    return pd.isna(value) or str(value).strip().lower() in ('','nan','none','null','-')


def clean_label(value):
    if blank(value): return ''
    val = str(value).strip()
    return val[:-2] if val.endswith('.0') and val[:-2].isdigit() else val


def normalize_shift(value):
    if blank(value): return None
    token = str(value).strip().upper()
    if token not in ALIASES:
        raise ValueError(f'인식할 수 없는 근무코드 {value!r}. 코드 대응표에 추가하거나 템플릿을 수정하세요.')
    return ALIASES[token]


def align_headers(df):
    out = df.copy()
    out.columns = [clean_label(c) for c in out.columns]
    if '그룹' not in out.columns or '이름' not in out.columns:
        for idx in range(min(5, len(out))):
            values = [clean_label(x) for x in out.iloc[idx]]
            if '그룹' in values and '이름' in values:
                out.columns = values
                out = out.iloc[idx+1:].reset_index(drop=True)
                break
    if '이름' not in out or '그룹' not in out:
        raise ValueError('템플릿에는 이름과 그룹 열이 필요합니다.')
    if len(set(out.columns)) != len(out.columns):
        raise ValueError('중복 또는 빈 헤더가 여러 개 있습니다. 각 열 이름을 고유하게 지정하세요.')
    return out


def read_table(name, data):
    df = pd.read_excel(io.BytesIO(data), dtype=object) if name.lower().endswith('.xlsx') else pd.read_csv(io.BytesIO(data), encoding='utf-8-sig', dtype=object)
    return align_headers(df)


def allowed_codes(value):
    if blank(value) or str(value).strip() == '전체': return set(SHIFTS)
    result = {normalize_shift(t) for t in re.split(r'[,/\s]+', str(value).strip()) if t}
    result.add('OFF')
    return result


def parse_days(value, days, title):
    if blank(value): return set()
    result = set()
    for token in re.split(r'[,/\s]+', str(value).strip()):
        if not re.fullmatch(r'\d+(?:\.0)?', token): raise ValueError(f'{title}: 날짜는 쉼표로 구분한 정수로 입력하세요.')
        day = int(float(token))
        if not 1 <= day <= days: raise ValueError(f'{title}: {day}일은 해당 월의 날짜가 아닙니다.')
        result.add(day)
    return result


def issue(kind, name, when, rule, detail):
    return dict(zip(ISSUE_COLUMNS, (kind, name, str(when), rule, detail)))


def parse_problem(df, year, month, previous=None):
    df = align_headers(df).reset_index(drop=True)
    days = calendar.monthrange(year, month)[1]
    day_cols = [c for c in df.columns if c.isdigit()]
    if set(day_cols) != {str(d) for d in range(1, days+1)}:
        raise ValueError(f'{year}년 {month}월에는 1~{days} 날짜 열이 연속으로 있어야 합니다. 월 선택과 날짜 열을 확인하세요.')
    start = next((i for i, row in df.iterrows() if any(('듀티별' in str(v) or '인원수' in str(v)) for v in row)), None)
    if start is None: raise ValueError('듀티별 인원수 행이 없습니다. D E N (선택 DE) 필요인원을 명시해 주세요.')
    req = {}
    de_groups = None
    for _, row in df.iloc[start:].iterrows():
        vals = [str(row[c]).strip().upper() for c in df.columns if not c.isdigit() and not blank(row[c])]
        duty = next((v for v in vals if v in DUTIES), None)
        if duty is None: continue
        if duty in req: raise ValueError(f'{duty} 필요인원 행이 중복되어 있습니다.')
        numbers=[]
        for d in range(1, days+1):
            value=row[str(d)]
            try: f=float(value)
            except (TypeError,ValueError): raise ValueError(f'{d}일 {duty} 필요인원이 비어 있거나 숫자가 아닙니다.')
            if not math.isfinite(f) or f < 0 or f != int(f): raise ValueError(f'{d}일 {duty} 필요인원은 0 이상의 정수여야 합니다.')
            numbers.append(int(f))
        req[duty]=numbers
        if duty=='DE' and not blank(row['그룹']):
            text=str(row['그룹']).strip().upper()
            if '듀티별' not in text and '인원수' not in text and text!='DE':
                de_groups=set(re.split(r'[,/\s]+',text))
    if not all(s in req for s in ('D','E','N')): raise ValueError('D E N 필요인원 행을 모두 입력하세요.')
    req.setdefault('DE',[0]*days)
    names=[];groups=[];keepers=[];allowed=[];wanted=[];rows=[];fixed_rows=[];codes=[]
    prev_map={}; warnings=[]
    if previous is not None:
        prev=align_headers(previous)
        prev_year, prev_month=(year-1,12) if month==1 else (year,month-1)
        last=calendar.monthrange(prev_year,prev_month)[1]
        expected=[str(d) for d in range(last-6,last+1)]
        if any(c not in prev.columns for c in expected):
            raise ValueError(f'이전 달 파일에는 {prev_year}년 {prev_month}월 마지막 7일 날짜 열이 필요합니다.')
        for _,row in prev.iterrows():
            name=clean_label(row['이름'])
            if not name or name.upper() in DUTIES or '듀티별' in str(row['그룹']) or '인원수' in str(row['그룹']):continue
            if name in prev_map: raise ValueError(f'이전 달 직원 이름 또는 식별자가 중복됩니다: {name}')
            if any(blank(row[c]) for c in expected): raise ValueError(f'{name}: 이전 달 마지막 7일 근무가 비어 있습니다.')
            prev_map[name]=[normalize_shift(row[c]) for c in expected]
    previous_group=''
    for idx,row in df.iloc[:start].iterrows():
        name=clean_label(row['이름'])
        if not name or name in ('이름','요일','월','화','수','목','금','토','일'):continue
        if name in names: raise ValueError(f'직원 이름 또는 식별자가 중복됩니다: {name}')
        group=clean_label(row['그룹']) or previous_group
        if not group: raise ValueError(f'{name}: 그룹을 입력하세요.')
        previous_group=group
        keeper='야간전담' in name or '야간전담' in group
        col='가능 근무' if '가능 근무' in df.columns else '가능근무'
        allow=allowed_codes(row[col]) if col in df.columns else set(SHIFTS)
        if keeper: allow={'N','OFF'}
        if de_groups is not None and group.upper() not in de_groups:allow.discard('DE')
        wants=parse_days(row.get('원티드 오프'),days,f'{name} 원티드 오프')
        approved=parse_days(row.get('승인 OFF'),days,f'{name} 승인 OFF')
        day_codes=[normalize_shift(row[str(d)]) for d in range(1,days+1)]
        for d in approved:
            if day_codes[d-1] not in (None,'OFF'):raise ValueError(f'{name} {d}일: 승인 OFF와 기입 근무가 충돌합니다.')
            day_codes[d-1]='OFF'
        names.append(name);groups.append(group);keepers.append(keeper);allowed.append(allow);wanted.append(wants);rows.append(idx)
        codes.append(day_codes);fixed_rows.append([x is not None for x in day_codes])
    if not names: raise ValueError('근무표에 직원이 없습니다.')
    histories=[]
    for name in names:
        history=prev_map.get(name,[])
        if not history:warnings.append(issue('확인 필요',name,'월초','이전 달 이력 없음','월초 경계를 완전히 검수하지 못했습니다. 이전 달 근무표를 업로드하세요.'))
        histories.append(history)
    return Problem(df,year,month,names,groups,keepers,allowed,wanted,np.array(fixed_rows,bool),np.array(codes,object),histories,warnings,req,rows)


def counted_code(code,cfg):
    return 'D' if code=='교육' and cfg.education_counts_as_day else code


def initialize_schedule(problem,cfg,rng):
    """고정 셀을 변경하지 않고 가능근무를 반영한 날짜별 이분 매칭. 불가능하면 중단."""
    n=len(problem.names); result=np.full((n,problem.days),'OFF',object)
    for d in range(problem.days):
        counts={s:0 for s in DUTIES}; free=[]
        for i in range(n):
            if problem.fixed[i,d]:
                code=problem.fixed_codes[i,d]
                if code not in problem.allowed[i]:raise ValueError(f'{problem.names[i]} {d+1}일: 고정 {code}와 가능근무가 충돌합니다. 고정 근무는 해제하지 않았습니다.')
                result[i,d]=code
                normalized=counted_code(code,cfg)
                if normalized in counts:counts[normalized]+=1
            else:free.append(i)
        for s in DUTIES:
            if counts[s]>problem.requirements[s][d]:raise ValueError(f'{d+1}일 {s}: 고정 인원 {counts[s]}명이 필요인원 {problem.requirements[s][d]}명을 초과합니다.')
        slots=[s for s in DUTIES for _ in range(problem.requirements[s][d]-counts[s])]
        if len(slots)>len(free):raise ValueError(f'{d+1}일: 미고정 직원 {len(free)}명으로 잔여 {len(slots)}개 근무를 채울 수 없습니다. 고정 OFF와 승인 OFF를 유지하고 생성을 중단했습니다.')
        # 제한적인 근무부터 매칭하고 재귀적 경로로 기존 배치를 재연결합니다.
        candidates={s:[i for i in free if s in problem.allowed[i]] for s in DUTIES}
        for s in DUTIES:rng.shuffle(candidates[s])
        # N에서 야간전담을 우선 고려하되 필요 시 일반직으로 매칭합니다.
        candidates['N'].sort(key=lambda i:not problem.keepers[i])
        order=sorted(range(len(slots)),key=lambda j:len(candidates[slots[j]]))
        assigned={}
        def match(j,seen):
            for i in candidates[slots[j]]:
                if i in seen:continue
                seen.add(i)
                if i not in assigned or match(assigned[i],seen):
                    assigned[i]=j;return True
            return False
        for j in order:
            if not match(j,set()):raise ValueError(f'{d+1}일 {slots[j]}: 가능근무와 고정근무를 유지하며 필요한 인원을 배치할 수 없습니다.')
        for i,j in assigned.items():result[i,d]=slots[j]
    return result


def row_issues(row,i,problem,cfg):
    """최적화 여부와 독립적으로 직원별 필수·희망조건을 항목별 검사."""
    out=[]; name=problem.names[i];history=problem.histories[i];h=len(history)
    raw=list(history)+list(row);norm=['D' if s=='교육' else s for s in raw]
    total=len(norm)
    def add(kind,d,rule,detail):
        when=date(problem.year,problem.month,d-h+1).isoformat() if h<=d<total else '월 경계'
        out.append(issue(kind,name,when,rule,detail))
    for k,code in enumerate(row):
        d=h+k
        if code not in SHIFTS:add('필수 위반',d,'근무코드',f'알 수 없는 코드 {code}')
        if code not in problem.allowed[i]:add('필수 위반',d,'가능근무',f'{code}는 허용 근무가 아닙니다.')
        if problem.fixed[i,k] and code!=problem.fixed_codes[i,k]:add('필수 위반',d,'고정근무 보호',f'고정 {problem.fixed_codes[i,k]}가 {code}로 바뀌었습니다.')
        if k+1 in problem.wanted[i] and code!='OFF':add('희망 미반영',d,'원티드 오프',f'희망 OFF 대신 {code} 배치')
    total_n=sum(s=='N' for s in row)
    if problem.keepers[i]:
        if total_n!=cfg.night_keeper_target:out.append(issue('필수 위반',name,'월 전체','야간전담 N 횟수',f'목표 {cfg.night_keeper_target}회 / 실제 {total_n}회'))
    elif total_n>cfg.max_night:out.append(issue('필수 위반',name,'월 전체','월간 N 상한',f'상한 {cfg.max_night}회 / 실제 {total_n}회'))
    cw=cn=0
    for d,s in enumerate(norm):
        cw=cw+1 if s!='OFF' else 0
        cn=cn+1 if s=='N' else 0
        if d>=h and cfg.max_work and cw>cfg.max_work:add('필수 위반',d,'최대 연속근무',f'{cw}일 연속 / 상한 {cfg.max_work}일')
        if d>=h and cn>cfg.max_consecutive_night:add('필수 위반',d,'최대 연속 N',f'{cn}일 연속 N / 상한 {cfg.max_consecutive_night}일')
        if d+1<total and d+1>=h:
            nxt=norm[d+1]
            if (s=='E' and nxt in ('D','DE')) or (s=='N' and nxt in ('D','E','DE')) or (s=='DE' and nxt=='D'):
                add('필수 위반',d+1,'금지 근무조합',f'{s} → {nxt}')
        if d+2<total and d+2>=h:
            pattern=tuple(norm[d:d+3])
            if pattern in [('E','OFF','D'),('N','OFF','D')] or (not problem.keepers[i] and pattern==('N','OFF','N')):
                add('필수 위반',d+2,'금지 3일 조합',' → '.join(pattern))
        if d+4<total and d+4>=h and tuple(norm[d:d+5]) in FORBIDDEN:
            add('필수 위반',d+4,'설정된 5일 금지패턴',' → '.join(norm[d:d+5]))
        # 연속 5일 도달 시 두 OFF를 검사. 모르는 다음 달은 확인 필요로 분리.
        if cfg.two_off_after_five and cw==5:
            for target in (d+1,d+2):
                if target<h:continue
                if target<total:
                    if norm[target]!='OFF':add('필수 위반',target,'5일 근무 후 2 OFF','5일 근무 다음 두 날짜에는 OFF가 필요합니다.')
                elif d>=h:add('확인 필요',d,'월말 후속 OFF','다음 달 근무표에서 5일 근무 후 OFF를 확인하세요.')
        if cfg.two_off_after_night and s=='N' and (d==total-1 or norm[d+1]!='N'):
            for target in (d+1,d+2):
                if target<h:continue
                if target<total:
                    if norm[target]!='OFF':add('필수 위반',target,'N 종료 후 2 OFF','N 종료 다음 두 날짜에는 OFF가 필요합니다.')
                elif d>=h:add('확인 필요',d,'월말 N 후속 OFF','다음 달 N 연속 여부와 종료 후 2 OFF를 확인하세요.')
        if d>=h and s=='N' and cfg.no_single_night:
            prev=d>0 and norm[d-1]=='N';nxt=d+1<total and norm[d+1]=='N'
            if not prev and not nxt:
                if d==0 or d==total-1:add('확인 필요',d,'경계 단독 N','인접 월 근무를 확인해야 단독 N 여부를 확정할 수 있습니다.')
                else:add('필수 위반',d,'단독 N','하루짜리 N 근무')
        if d>=h and s!='OFF' and cfg.no_single_work:
            if d>0 and d+1<total and norm[d-1]=='OFF' and norm[d+1]=='OFF':add('희망 미반영',d,'단독 근무','OFF 사이 하루짜리 근무')
            elif (d==0 or d==total-1) and ((d==0 and total>1 and norm[1]=='OFF') or (d==total-1 and d>0 and norm[d-1]=='OFF')):
                add('확인 필요',d,'경계 단독근무','인접 월 근무 확인 필요')
    # 동일 항목의 동일 날짜 중복을 제거합니다.
    return [dict(t) for t in dict.fromkeys(tuple(x.items()) for x in out)]


def day_group_penalty(col,problem,cfg):
    if not cfg.group_balance:return 0
    groups=sorted(set(problem.groups));penalty=0
    for s in ('D','E','N','OFF'):
        counts=[sum(col[i]==s and problem.groups[i]==g for i in range(len(col))) for g in groups]
        total=sum(counts);low=total//len(groups);high=math.ceil(total/len(groups))
        penalty+=sum(max(0,low-c,c-high) for c in counts)*50
    return penalty


def row_score(row,i,problem,cfg):
    errors=row_issues(row,i,problem,cfg)
    hard=sum(e['구분']=='필수 위반' for e in errors)
    soft=sum(10000 if e['항목']=='원티드 오프' else 1000 for e in errors if e['구분']=='희망 미반영')
    if not problem.keepers[i]:
        normals=sum(not k for k in problem.keepers)
        avg_n=max(0,sum(problem.requirements['N'])-sum(problem.keepers)*cfg.night_keeper_target)/max(1,normals)
        avg_off=max(0,normals*problem.days-sum(sum(v) for v in problem.requirements.values())+sum(problem.keepers)*cfg.night_keeper_target)/max(1,normals)
        soft+=(abs(sum(s=='N' for s in row)-avg_n)*100+abs(sum(s=='OFF' for s in row)-avg_off)*80)
    return hard,soft


def validate_schedule(schedule,problem,cfg):
    if schedule.shape!=(len(problem.names),problem.days):raise ValueError('근무표 배열 크기가 입력과 일치하지 않습니다.')
    records=list(problem.warnings)
    for d in range(problem.days):
        for s in DUTIES:
            actual=sum(counted_code(x,cfg)==s for x in schedule[:,d]); required=problem.requirements[s][d]
            if actual!=required:records.append(issue('필수 위반','부서 전체',problem.dates[d].isoformat(),'필요인원',f'{s} 필요 {required}명 / 실제 {actual}명'))
    for i in range(len(problem.names)):records.extend(row_issues(schedule[i],i,problem,cfg))
    return pd.DataFrame(records,columns=ISSUE_COLUMNS).drop_duplicates(ignore_index=True)


def result_status(issues):
    if any(issues['구분']=='필수 위반'):return '검토 필요 — 필수조건 미충족'
    if any(issues['구분']=='확인 필요'):return '검토 필요 — 월 경계 등 확인 필요'
    return '조건 충족 — 검수 범위 내'


def distribution(schedule,problem):
    records=[];weekend=[d for d,dt in enumerate(problem.dates) if dt.weekday()>=5]
    for i,name in enumerate(problem.names):
        row=schedule[i]
        records.append({'이름':name,'그룹':problem.groups[i],'야간전담':problem.keepers[i],
                        **{s:int(sum(row==s)) for s in SHIFTS},
                        '총 근무':int(sum(row!='OFF')),
                        '주말근무':sum(row[d]!='OFF' for d in weekend),
                        '주말 OFF':sum(row[d]=='OFF' for d in weekend),
                        '원티드 요청':len(problem.wanted[i]),
                        '원티드 반영':sum(row[d-1]=='OFF' for d in problem.wanted[i])})
    return pd.DataFrame(records)


def optimize(problem,cfg,progress:Callable|None=None):
    if not 0<=cfg.max_work<=5 or not 0<=cfg.max_night<=31 or not 1<=cfg.max_consecutive_night<=31:raise ValueError('근무조건 범위가 잘못되었습니다.')
    rng=random.Random(cfg.seed);sched=initialize_schedule(problem,cfg,rng)
    scores=[row_score(sched[i],i,problem,cfg) for i in range(len(problem.names))]
    columns=[day_group_penalty(sched[:,d],problem,cfg) for d in range(problem.days)]
    hard=sum(x[0] for x in scores);soft=sum(x[1] for x in scores)+sum(columns)
    best=sched.copy();best_score=(hard,soft);current=(hard,soft)
    started=time.monotonic();attempts=0
    if len(problem.names)>1:
        for step in range(cfg.max_iterations):
            attempts=step+1
            if time.monotonic()-started>=cfg.time_limit:break
            if progress and step%1000==0:progress(step/cfg.max_iterations)
            # 동일 날짜 교환은 필요인원 구성을 보존합니다. 고정 셀은 후보에서 제외.
            d=rng.randrange(problem.days);free=np.flatnonzero(~problem.fixed[:,d]).tolist()
            if len(free)<2:continue
            a,b=rng.sample(free,2);sa,sb=sched[a,d],sched[b,d]
            if sa==sb or sb not in problem.allowed[a] or sa not in problem.allowed[b]:continue
            olda,oldb=scores[a],scores[b];oldcol=columns[d]
            sched[a,d],sched[b,d]=sb,sa
            na=row_score(sched[a],a,problem,cfg);nb=row_score(sched[b],b,problem,cfg);nc=day_group_penalty(sched[:,d],problem,cfg)
            newhard=current[0]-olda[0]-oldb[0]+na[0]+nb[0]
            newsoft=current[1]-olda[1]-oldb[1]-oldcol+na[1]+nb[1]+nc
            candidate=(newhard,newsoft)
            # 필수 위반 개수를 우선 비교하고, 같은 개수일 때 선호 점수를 탐색합니다.
            # 탈출 탐색은 소량의 악화를 확률적으로 허용하되 최저점 결과는 별도 보존.
            temperature=max(.1,4*(1-step/max(1,cfg.max_iterations)))
            delta=(newhard-current[0])*10+(newsoft-current[1])/10000
            accepted=candidate<current or rng.random()<math.exp(-max(0,delta)/temperature)
            if accepted:
                scores[a],scores[b],columns[d]=na,nb,nc;current=candidate
                if current<best_score:best=sched.copy();best_score=current
            else:sched[a,d],sched[b,d]=sa,sb
    issues=validate_schedule(best,problem,cfg)
    result_df=problem.source.copy()
    for i,idx in enumerate(problem.row_indices):
        for d in range(problem.days):result_df.at[idx,str(d+1)]=best[i,d]
    dist=distribution(best,problem)
    return {'schedule':best,'table':result_df,'issues':issues,'distribution':dist,
            'status':result_status(issues),'settings':asdict(cfg),'elapsed':time.monotonic()-started,
            'iterations':attempts,'score':best_score,'month':f'{problem.year}-{problem.month:02d}'}


def export_result(result):
    """결과 상태와 검수내역을 항상 근무표와 함께 내보냅니다."""
    data=io.BytesIO()
    info={'대상월':result['month'],'판정':result['status'],'실행시간 초':round(result['elapsed'],3),
          '탐색횟수':result['iterations'],'필수 위반 수':int(sum(result['issues']['구분']=='필수 위반')),
          '사용방법':'관리자 확인 후 확정. 조건 충족은 설정된 검사 범위만 의미합니다.',
          **result['settings']}
    with pd.ExcelWriter(data,engine='openpyxl') as writer:
        result['table'].to_excel(writer,sheet_name='근무표',index=False)
        result['issues'].to_excel(writer,sheet_name='검수내역',index=False)
        result['distribution'].to_excel(writer,sheet_name='개인별 배분',index=False)
        pd.DataFrame(info.items(),columns=['항목','값']).to_excel(writer,sheet_name='생성 정보',index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='C2' if ws.title=='근무표' else 'A2'
            ws.auto_filter.ref=ws.dimensions
            for cells in ws.columns:
                letter=cells[0].column_letter
                ws.column_dimensions[letter].width=min(50,max(10,max(len(str(c.value or '')) for c in cells)+3))
            for cell in ws[1]:
                from openpyxl.styles import Font,PatternFill
                cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24445C')
            # 업로드한 텍스트가 수식으로 실행되지 않도록 모든 문자열 셀을 텍스트로 저장.
            for cells in ws.iter_rows():
                for cell in cells:
                    if isinstance(cell.value,str):cell.data_type='s'
    return data.getvalue()


def example_template(year=2026,month=10):
    days=calendar.monthrange(year,month)[1];records=[]
    for i in range(12):
        records.append({'이름':f'예시{i+1:02d}','그룹':('A','B','C')[i%3],
                        '가능 근무':'D,E,N','원티드 오프':'','승인 OFF':'',
                        **{str(d):'' for d in range(1,days+1)}})
    for j,s in enumerate(DUTIES):
        records.append({'이름':s,'그룹':'듀티별 인원수' if j==0 else '',
                        '가능 근무':'','원티드 오프':'','승인 OFF':'',
                        **{str(d):2 if s in ('D','E') else 1 if s=='N' else 0 for d in range(1,days+1)}})
    return pd.DataFrame(records)


def app():
    import streamlit as st
    st.set_page_config(page_title='스마트 널스 스케줄러 v2',layout='wide')
    st.title('스마트 널스 스케줄러 v2')
    st.caption('관리자용 근무표 초안 생성 · 고정근무 보호 · 자동 검수')
    st.info('템플릿의 기입 근무와 승인 OFF는 변경하지 않습니다. 원티드 오프는 희망조건입니다. 생성 결과를 검수한 뒤 확정하세요.')
    st.sidebar.header('대상 월과 파일')
    picked=st.sidebar.date_input('대상 월 선택',value=date.today().replace(day=1))
    year,month=picked.year,picked.month
    uploaded=st.sidebar.file_uploader('1 초기 근무표',type=['xlsx','csv'])
    previous=st.sidebar.file_uploader('2 이전 달 근무표',type=['xlsx','csv'])
    st.sidebar.header('부서별 근무조건')
    cfg=Settings(
        max_work=st.sidebar.slider('최대 연속근무 일수 0은 제한 끄기',0,5,5),
        max_night=st.sidebar.slider('일반직 월 N 상한',0,7,6),
        max_consecutive_night=st.sidebar.slider('최대 연속 N',2,3,3),
        night_keeper_target=st.sidebar.number_input('야간전담 월 N 목표',min_value=0,max_value=31,value=15),
        two_off_after_five=st.sidebar.toggle('5일 근무 후 2 OFF 검사',value=True),
        no_single_night=st.sidebar.toggle('단독 N 금지 검사',value=True),
        group_balance=st.sidebar.toggle('날짜별 그룹 배분 평가',value=True),
        two_off_after_night=st.sidebar.toggle('N 종료 후 2 OFF 검사',value=True),
        no_single_work=st.sidebar.toggle('단독 근무 줄이기 희망조건',value=True),
        education_counts_as_day=st.sidebar.toggle('교육을 D 필요인원에 포함',value=False),
        max_iterations=st.sidebar.slider('최대 탐색 횟수',10000,150000,60000,step=10000),
        time_limit=float(st.sidebar.slider('최대 탐색시간 초',5,120,30)),
        seed=st.sidebar.number_input('탐색 번호 같은 입력은 동일 시작점',min_value=0,value=42),
    )
    st.sidebar.caption('야간 후 휴무·금지패턴은 부서 운영조건입니다. 법정 근로시간을 계산하는 기능은 아닙니다.')
    demo=io.BytesIO();example_template(year,month).to_excel(demo,index=False)
    st.download_button('해당 월 예시 템플릿 다운로드',demo.getvalue(),f'예시_템플릿_{year}_{month:02d}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with st.expander('입력 및 검수 기준'):
        st.write('이름, 그룹, 가능 근무, 원티드 오프, 승인 OFF와 날짜 열을 사용합니다. 하단에 D E N 필요인원 행을 명시하세요. DE는 선택입니다. 기입 근무는 야간전담을 포함해 모두 고정됩니다.')
        st.write('원티드는 미반영될 수 있으며 내역이 표시됩니다. 승인 휴무는 승인 OFF 열이나 날짜 칸 OFF로 입력합니다. 교육은 가능근무에서 교육으로 허용해야 하며 D 인원 포함 여부를 선택할 수 있습니다.')
        st.write('주말은 토요일·일요일입니다. 주말 배분 현황은 표시하며 자동 균등화는 하지 않습니다. 인접 월 이력이 없거나 월말 후속 조건을 확인할 수 없으면 검토 필요로 표시됩니다.')
        st.write('로그인·권한 관리와 다중 사용자 신청 저장은 포함하지 않은 관리자 단독 사용 버전입니다. 로컬 실행은 127.0.0.1 주소를 사용하세요.')
    if not uploaded:
        st.info('왼쪽에서 초기 근무표를 업로드하세요.');return
    try:
        data=uploaded.getvalue();prevdata=previous.getvalue() if previous else b''
        base=read_table(uploaded.name,data)
    except Exception as exc:
        st.error(f'파일 읽기 오류: {exc}');return
    signature=hashlib.sha256(data+prevdata+str((year,month,asdict(cfg))).encode()).hexdigest()
    if st.session_state.get('input_signature')!=signature:
        st.session_state['input_signature']=signature;st.session_state['result']=None
        st.session_state['editor_revision']=st.session_state.get('editor_revision',0)+1
    st.subheader('원티드 및 승인 OFF 확인')
    st.caption('날짜별 고정근무 수정은 원본 템플릿에서 하세요. 여기서는 원티드 오프와 승인 OFF만 수정할 수 있습니다.')
    for c in ('원티드 오프','승인 OFF'):
        if c not in base:base[c]=''
        base[c]=base[c].fillna('').astype(str)
    editable=st.data_editor(base,disabled=[c for c in base.columns if c not in ('원티드 오프','승인 OFF')],hide_index=True,key=f"editor_{st.session_state['editor_revision']}")
    current_edit=hashlib.sha256(editable.to_csv(index=False).encode()).hexdigest()
    result=st.session_state.get('result')
    if result and st.session_state.get('result_edit')!=current_edit:
        st.session_state['result']=None;result=None
    if st.button('근무표 초안 생성 및 검수',type='primary'):
        st.session_state['result']=None
        bar=st.progress(0.0)
        try:
            prev=read_table(previous.name,prevdata) if previous else None
            problem=parse_problem(editable,year,month,prev)
            result=optimize(problem,cfg,lambda x:bar.progress(min(.99,x)))
            st.session_state['result']=result;st.session_state['result_edit']=current_edit
            bar.progress(1.0)
        except Exception as exc:
            bar.empty();st.error(f'생성을 중단했습니다: {exc}');return
    result=st.session_state.get('result')
    if result is None:return
    issues=result['issues'];status=result['status']
    if status.startswith('조건 충족'):st.success(status)
    else:st.warning(status)
    cols=st.columns(3)
    cols[0].metric('필수 위반',int(sum(issues['구분']=='필수 위반')))
    cols[1].metric('희망 미반영',int(sum(issues['구분']=='희망 미반영')))
    cols[2].metric('확인 필요',int(sum(issues['구분']=='확인 필요')))
    st.caption(f"탐색시간 {result['elapsed']:.1f}초 · 탐색 시도 {result['iterations']:,}회 · 관리자가 확인한 뒤 확정하세요.")
    tab1,tab2,tab3=st.tabs(['생성 근무표','검수 내역','개인별 근무 배분'])
    with tab1:st.dataframe(result['table'],hide_index=True)
    with tab2:
        if issues.empty:st.write('설정된 검사 범위에서 위반이 발견되지 않았습니다.')
        else:st.dataframe(issues,hide_index=True)
    with tab3:
        st.dataframe(result['distribution'],hide_index=True)
        st.caption('주말은 토요일과 일요일이며 교육과 DE도 근무일로 집계합니다. 개인별 가능근무가 다르면 횟수 차이만으로 형평성을 판단하지 마세요.')
    prefix='검토필요' if status.startswith('검토 필요') else '조건충족'
    st.download_button('검수내역 포함 엑셀 다운로드',export_result(result),f'{prefix}_근무표_{year}_{month:02d}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__=='__main__':app()
